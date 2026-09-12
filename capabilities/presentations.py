import os

from capabilities.base import Cap, fail, ok, register_cap, run_shell

try:
    import pptx
except ImportError:
    pptx = None

def install(registry, *, approve_all=False):
    @register_cap(registry, "pres.create", Cap.WRITE, "Create new presentation", approve_all)
    def create(path: str):
        if pptx and path.endswith(".pptx"):
            prs = pptx.Presentation()
            prs.save(path)
            return ok(f"Created {path}")
        return run_shell(["touch", path])

    @register_cap(registry, "pres.open", Cap.LOW, "Open in LibreOffice Impress", approve_all)
    def open_pres(path: str):
        return run_shell(["libreoffice", "--impress", path])

    @register_cap(registry, "pres.add_slide", Cap.WRITE, "Add slide", approve_all)
    def add_slide(path: str, index: int, layout: int):
        if pptx and path.endswith(".pptx"):
            try:
                prs = pptx.Presentation(path)
                slide_layout = prs.slide_layouts[layout]
                prs.slides.add_slide(slide_layout)
                prs.save(path)
                return ok(f"Added slide to {path}")
            except Exception as e:
                return fail(str(e))
        return fail("Unsupported format")

    @register_cap(registry, "pres.delete_slide", Cap.WRITE, "Delete slide", approve_all)
    def delete_slide(path: str, index: int):
        return fail("Not easily supported by python-pptx natively")

    @register_cap(registry, "pres.add_text", Cap.WRITE, "Add text to slide", approve_all)
    def add_text(path: str, slide_index: int, text: str, placeholder: int):
        if pptx and path.endswith(".pptx"):
            try:
                prs = pptx.Presentation(path)
                slide = prs.slides[slide_index]
                shape = slide.placeholders[placeholder]
                shape.text = text
                prs.save(path)
                return ok(f"Added text to slide {slide_index}")
            except Exception as e:
                return fail(str(e))
        return fail("Unsupported format")

    @register_cap(registry, "pres.add_image", Cap.WRITE, "Add image to slide", approve_all)
    def add_image(path: str, slide_index: int, image_path: str, x: int, y: int, width: int, height: int):
        if pptx and path.endswith(".pptx"):
            try:
                prs = pptx.Presentation(path)
                slide = prs.slides[slide_index]
                from pptx.util import Inches
                slide.shapes.add_picture(image_path, Inches(x), Inches(y), Inches(width), Inches(height))
                prs.save(path)
                return ok(f"Added image to slide {slide_index}")
            except Exception as e:
                return fail(str(e))
        return fail("Unsupported format")

    @register_cap(registry, "pres.get_slide_count", Cap.READ, "Get slide count", approve_all)
    def get_slide_count(path: str):
        if pptx and path.endswith(".pptx"):
            try:
                prs = pptx.Presentation(path)
                return ok(str(len(prs.slides)))
            except Exception as e:
                return fail(str(e))
        return fail("Unsupported format")

    @register_cap(registry, "pres.export", Cap.WRITE, "Export to PDF", approve_all)
    def export(path: str, output_path: str):
        outdir = os.path.dirname(output_path) or "."
        return run_shell(["libreoffice", "--headless", "--convert-to", "pdf", path, "--outdir", outdir])

    @register_cap(registry, "pres.save", Cap.WRITE, "Save presentation", approve_all)
    def save(path: str):
        return ok(f"Saved {path}")
